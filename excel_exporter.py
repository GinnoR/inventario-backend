"""
Generador de Excel profesional para Inventario Por Imágenes.
Crea un .xlsx con imágenes de categoría incrustadas en cada fila,
alturas de fila apropiadas, y formato profesional.
"""

import io
import os
import tempfile
from typing import List, Dict, Any
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle
)
from openpyxl.utils import get_column_letter
from datetime import datetime

# ─── Category config ───────────────────────────────────
CATEGORY_CONFIG = {
    'abarrotes':        {'emoji': '🛒', 'color': '#F59E0B', 'bg': '#FEF3C7', 'icon': 'AB'},
    'bebidas':          {'emoji': '🥤', 'color': '#06B6D4', 'bg': '#CFFAFE', 'icon': 'BE'},
    'lácteos':          {'emoji': '🥛', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'LA'},
    'lacteos':          {'emoji': '🥛', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'LA'},
    'limpieza':         {'emoji': '🧹', 'color': '#14B8A6', 'bg': '#CCFBF1', 'icon': 'LI'},
    'snacks':           {'emoji': '🍿', 'color': '#F97316', 'bg': '#FED7AA', 'icon': 'SN'},
    'galletas':         {'emoji': '🍪', 'color': '#CA8A04', 'bg': '#FEF9C3', 'icon': 'GA'},
    'enlatados':        {'emoji': '🥫', 'color': '#EF4444', 'bg': '#FEE2E2', 'icon': 'EN'},
    'conservas':        {'emoji': '🥫', 'color': '#EF4444', 'bg': '#FEE2E2', 'icon': 'CO'},
    'condimentos':      {'emoji': '🧂', 'color': '#EAB308', 'bg': '#FEF9C3', 'icon': 'CD'},
    'salsas':           {'emoji': '🫙', 'color': '#F43F5E', 'bg': '#FFE4E6', 'icon': 'SA'},
    'cuidado personal': {'emoji': '🧴', 'color': '#EC4899', 'bg': '#FCE7F3', 'icon': 'CP'},
    'higiene':          {'emoji': '🧼', 'color': '#8B5CF6', 'bg': '#EDE9FE', 'icon': 'HI'},
    'frutas':           {'emoji': '🍎', 'color': '#22C55E', 'bg': '#DCFCE7', 'icon': 'FR'},
    'verduras':         {'emoji': '🥬', 'color': '#10B981', 'bg': '#D1FAE5', 'icon': 'VE'},
    'carnes':           {'emoji': '🥩', 'color': '#DC2626', 'bg': '#FEE2E2', 'icon': 'CA'},
    'panadería':        {'emoji': '🍞', 'color': '#D97706', 'bg': '#FDE68A', 'icon': 'PA'},
    'cereales':         {'emoji': '🥣', 'color': '#FBBF24', 'bg': '#FEF3C7', 'icon': 'CE'},
    'aceites':          {'emoji': '🫒', 'color': '#84CC16', 'bg': '#ECFCCB', 'icon': 'AC'},
    'pastas':           {'emoji': '🍝', 'color': '#EAB308', 'bg': '#FEF9C3', 'icon': 'PS'},
    'mascotas':         {'emoji': '🐕', 'color': '#8B5CF6', 'bg': '#EDE9FE', 'icon': 'MA'},
    'embutidos':        {'emoji': '🌭', 'color': '#F97316', 'bg': '#FFEDD5', 'icon': 'EM'},
    'herramientas':     {'emoji': '🔨', 'color': '#64748B', 'bg': '#F1F5F9', 'icon': 'HE'},
    'ferretería':       {'emoji': '🔩', 'color': '#475569', 'bg': '#E2E8F0', 'icon': 'FE'},
    'ferreteria':       {'emoji': '🔩', 'color': '#475569', 'bg': '#E2E8F0', 'icon': 'FE'},
    'electricidad':     {'emoji': '🔌', 'color': '#EAB308', 'bg': '#FEF9C3', 'icon': 'EL'},
    'línea sanitaria':  {'emoji': '🚰', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'LS'},
    'linea sanitaria':  {'emoji': '🚰', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'LS'},
    'fontanería':       {'emoji': '🚰', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'FO'},
    'fontaneria':       {'emoji': '🚰', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'FO'},
    'plomería':         {'emoji': '🚰', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'PL'},
    'plomeria':         {'emoji': '🚰', 'color': '#3B82F6', 'bg': '#DBEAFE', 'icon': 'PL'},
    'pinturas':         {'emoji': '🎨', 'color': '#8B5CF6', 'bg': '#EDE9FE', 'icon': 'PI'},
    'construcción':     {'emoji': '🧱', 'color': '#D97706', 'bg': '#FEF3C7', 'icon': 'CT'},
    'jardinería':       {'emoji': '🪴', 'color': '#22C55E', 'bg': '#DCFCE7', 'icon': 'JA'},
}

CONF_COLORS = {
    'alta':  {'fill': '22C55E', 'font': 'FFFFFF'},
    'media': {'fill': 'F59E0B', 'font': 'FFFFFF'},
    'baja':  {'fill': 'EF4444', 'font': 'FFFFFF'},
}

def _get_cat_config(category: str) -> dict:
    key = category.lower().strip()
    for k, v in CATEGORY_CONFIG.items():
        if k in key:
            return v
    return {'emoji': '📦', 'color': '#64748B', 'bg': '#F1F5F9', 'icon': '??'}


def _hex_to_rgb(hex_color: str) -> tuple:
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def _create_category_image(category: str, size: int = 60) -> str:
    """
    Creates a small colored category badge image.
    Returns the path to a temporary PNG file.
    """
    config = _get_cat_config(category)
    bg_rgb = _hex_to_rgb(config['bg'])
    fg_rgb = _hex_to_rgb(config['color'])
    
    # Create image with rounded rectangle
    img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    
    # Draw rounded rectangle background
    draw.rounded_rectangle(
        [(2, 2), (size - 2, size - 2)],
        radius=12,
        fill=bg_rgb + (255,),
        outline=fg_rgb + (200,),
        width=2
    )
    
    # Draw category abbreviation text
    try:
        font = ImageFont.truetype("arial.ttf", size // 3)
    except (OSError, IOError):
        try:
            font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", size // 3)
        except (OSError, IOError):
            font = ImageFont.load_default()
    
    text = config['icon']
    bbox = draw.textbbox((0, 0), text, font=font)
    text_w = bbox[2] - bbox[0]
    text_h = bbox[3] - bbox[1]
    x = (size - text_w) // 2
    y = (size - text_h) // 2 - 2
    draw.text((x, y), text, fill=fg_rgb + (255,), font=font)
    
    # Save to temp file
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
    img.save(tmp.name, 'PNG')
    tmp.close()
    return tmp.name


def _safe_int(val: Any) -> int:
    try:
        if isinstance(val, str):
            import re
            m = re.search(r'\d+', val)
            return int(m.group()) if m else 0
        return int(val)
    except (ValueError, TypeError):
        return 0

def generate_inventory_excel(products: List[Dict[str, Any]]) -> io.BytesIO:
    """
    Generates a professional Excel file with embedded category images.
    Returns a BytesIO buffer containing the .xlsx file.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario Detectado"
    
    # ─── Styles ─────────────────────────────────────────
    header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    cell_font = Font(name='Calibri', size=10)
    cell_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )
    
    # ─── Title Row ──────────────────────────────────────
    ws.merge_cells('A1:M1')
    title_cell = ws['A1']
    title_cell.value = '📊 INVENTARIO DETECTADO POR IA'
    title_cell.font = Font(name='Calibri', bold=True, size=16, color='1E40AF')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    title_cell.fill = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')
    ws.row_dimensions[1].height = 40
    
    # ─── Subtitle Row ──────────────────────────────────
    ws.merge_cells('A2:M2')
    sub_cell = ws['A2']
    now = datetime.now().strftime('%d/%m/%Y %H:%M')
    total_units = sum(_safe_int(p.get('cantidad_estimada', 0)) for p in products)
    sub_cell.value = f'Generado: {now}  |  {len(products)} productos  |  {total_units} unidades totales  |  InventarioAI + Gemini'
    sub_cell.font = Font(name='Calibri', size=9, color='64748B', italic=True)
    sub_cell.alignment = Alignment(horizontal='center', vertical='center')
    sub_cell.fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    ws.row_dimensions[2].height = 22
    
    # ─── Empty spacer row ──────────────────────────────
    ws.row_dimensions[3].height = 8
    
    # ─── Headers ───────────────────────────────────────
    headers = ['#', 'Imagen', 'Producto', 'Marca', 'Categoría', 'Ubicación', 'Cant. Visible', 'Profundidad', 'Cant. Total', 'Rango', 'Método', 'Confianza', 'Observaciones']
    col_widths = [5, 12, 35, 18, 18, 20, 10, 10, 10, 10, 15, 12, 35]
    
    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    
    ws.row_dimensions[4].height = 28
    
    # ─── Data Rows ─────────────────────────────────────
    IMAGE_SIZE = 50  # pixels
    ROW_HEIGHT = 48  # points (appropriate for ~50px image)
    
    temp_files = []  # Track temp files for cleanup
    
    for idx, product in enumerate(products, 1):
        row = idx + 4  # offset by header rows
        
        # Row height for image
        ws.row_dimensions[row].height = ROW_HEIGHT
        
        # Alternating row colors
        row_fill = PatternFill(
            start_color='F8FAFC' if idx % 2 == 0 else 'FFFFFF',
            end_color='F8FAFC' if idx % 2 == 0 else 'FFFFFF',
            fill_type='solid'
        )
        
        # # Column
        cell = ws.cell(row=row, column=1, value=idx)
        cell.font = Font(name='Calibri', size=9, color='94A3B8')
        cell.alignment = center_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Image column (B) - Create and embed category image
        img_cell = ws.cell(row=row, column=2, value='')
        img_cell.fill = row_fill
        img_cell.border = thin_border
        img_cell.alignment = center_alignment
        
        try:
            image_url = product.get('image_url')
            if image_url:
                # Cargar imagen real desde el frontend public
                import os
                base_public = r"C:\Users\hp\.gemini\antigravity-ide\scratch\Inventario Por Imagenes\frontend\public"
                img_path = os.path.join(base_public, image_url.lstrip("/\\"))
                
                # Check if it exists and use it
                if os.path.exists(img_path):
                    # We might need to resize it to IMAGE_SIZE using Pillow
                    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
                    with Image.open(img_path) as im:
                        im = im.convert("RGBA")
                        im.thumbnail((IMAGE_SIZE, IMAGE_SIZE))
                        im.save(tmp.name, "PNG")
                    img_path = tmp.name
                    temp_files.append(img_path)
                else:
                    img_path = _create_category_image(product.get('categoria', ''), IMAGE_SIZE)
                    temp_files.append(img_path)
            else:
                img_path = _create_category_image(product.get('categoria', ''), IMAGE_SIZE)
                temp_files.append(img_path)
                
            xl_img = XlImage(img_path)
            xl_img.width = IMAGE_SIZE
            xl_img.height = IMAGE_SIZE
            # Anchor image in cell B{row}
            cell_ref = f'B{row}'
            ws.add_image(xl_img, cell_ref)
        except Exception as e:
            # Fallback: just put category emoji text
            cat_config = _get_cat_config(product.get('categoria', ''))
            img_cell.value = cat_config['emoji']
            img_cell.font = Font(size=20)
        
        # Product name (C)
        cell = ws.cell(row=row, column=3, value=product.get('nombre_producto', ''))
        cell.font = Font(name='Calibri', size=10, bold=True, color='1E293B')
        cell.alignment = cell_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Brand (D)
        cell = ws.cell(row=row, column=4, value=product.get('marca', ''))
        cell.font = Font(name='Calibri', size=10, color='475569')
        cell.alignment = cell_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        

        # Category (E)
        cat_config = _get_cat_config(product.get('categoria', ''))
        cell = ws.cell(row=row, column=5, value=product.get('categoria', ''))
        cell.font = Font(name='Calibri', size=10, color=cat_config['color'].lstrip('#'))
        cell.alignment = center_alignment
        cat_bg = cat_config['bg'].lstrip('#')
        cell.fill = PatternFill(start_color=cat_bg, end_color=cat_bg, fill_type='solid')
        cell.border = thin_border
        
        # Ubicacion (F)
        cell = ws.cell(row=row, column=6, value=product.get('ubicacion', ''))
        cell.font = Font(name='Calibri', size=10, color='475569')
        cell.alignment = center_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Cant. Visible (G)
        cell = ws.cell(row=row, column=7, value=str(product.get('cantidad_visible', '')))
        cell.font = Font(name='Calibri', size=10, color='1E293B')
        cell.alignment = center_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Profundidad (H)
        cell = ws.cell(row=row, column=8, value=str(product.get('profundidad_estimada', '')))
        cell.font = Font(name='Calibri', size=10, color='1E293B')
        cell.alignment = center_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Cantidad Total (I)
        qty = product.get('cantidad_estimada', '0')
        cell = ws.cell(row=row, column=9, value=int(qty) if str(qty).isdigit() else qty)
        cell.font = Font(name='Calibri', size=12, bold=True, color='1E293B')
        cell.alignment = center_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Rango (J)
        cell = ws.cell(row=row, column=10, value=str(product.get('rango_estimado', '')))
        cell.font = Font(name='Calibri', size=10, color='475569')
        cell.alignment = center_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Metodo (K)
        cell = ws.cell(row=row, column=11, value=product.get('metodo_conteo', ''))
        cell.font = Font(name='Calibri', size=10, color='475569')
        cell.alignment = center_alignment
        cell.fill = row_fill
        cell.border = thin_border
        
        # Confianza (L)
        conf = product.get('confianza', 'Baja')
        conf_key = conf.lower().strip()
        conf_style = CONF_COLORS.get(conf_key, CONF_COLORS['baja'])
        cell = ws.cell(row=row, column=12, value=conf)
        cell.font = Font(name='Calibri', size=10, bold=True, color=conf_style['font'])
        cell.fill = PatternFill(start_color=conf_style['fill'], end_color=conf_style['fill'], fill_type='solid')
        cell.alignment = center_alignment
        cell.border = thin_border
        
        # Observaciones (M)
        cell = ws.cell(row=row, column=13, value=product.get('observaciones', ''))
        cell.font = Font(name='Calibri', size=9, color='94A3B8')
        cell.alignment = cell_alignment
        cell.fill = row_fill
        cell.border = thin_border
    
    # ─── Summary Row ───────────────────────────────────
    summary_row = len(products) + 5
    ws.row_dimensions[summary_row].height = 30
    
    ws.merge_cells(f'A{summary_row}:D{summary_row}')
    cell = ws.cell(row=summary_row, column=1, value='TOTAL')
    cell.font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    cell.fill = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    cell.alignment = Alignment(horizontal='right', vertical='center')
    cell.border = thin_border
    
    # Empty styled cells for merged area
    for col in [2, 3, 4]:
        c = ws.cell(row=summary_row, column=col)
        c.fill = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
        c.border = thin_border
    
    # Category count
    cell = ws.cell(row=summary_row, column=5, 
                   value=f'{len(set(p.get("categoria", "") for p in products))} cats.')
    cell.font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    cell.fill = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    cell.alignment = center_alignment
    cell.border = thin_border
    
    # Total quantity
    cell = ws.cell(row=summary_row, column=9, value=total_units)
    cell.font = Font(name='Calibri', bold=True, size=12, color='FFFFFF')
    cell.fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
    cell.alignment = center_alignment
    cell.border = thin_border
    
    # Confidence summary
    high_conf = sum(1 for p in products if p.get('confianza', '').lower() == 'alta')
    cell = ws.cell(row=summary_row, column=12, value=f'{high_conf}/{len(products)} Alta')
    cell.font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    cell.fill = PatternFill(start_color='16A34A', end_color='16A34A', fill_type='solid')
    cell.alignment = center_alignment
    cell.border = thin_border
    
    # ─── Freeze panes ──────────────────────────────────
    ws.freeze_panes = 'A5'
    
    # ─── Print settings ────────────────────────────────
    ws.sheet_properties.pageSetUpPr = None
    ws.print_title_rows = '4:4'
    
    # ─── Save to buffer ────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    # Cleanup temp image files
    for f in temp_files:
        try:
            os.unlink(f)
        except OSError:
            pass
    
    return buffer
